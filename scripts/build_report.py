import json
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from src.dataset import discover_images


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "report" / "results"
OUTPUT = ROOT / "report" / "Assignment2_Avocado_Report.docx"


def add_bullet(document, text):
    document.add_paragraph(text, style="List Bullet")


def add_centered_picture(document, path, width=6.0):
    if path.is_file():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Inches(width))


def main():
    metrics_path = RESULTS / "model_comparison.json"
    eda_path = RESULTS / "eda" / "image_analysis.csv"
    if not metrics_path.is_file() or not eda_path.is_file():
        raise FileNotFoundError(
            "Run dataset inspection and model evaluation before building the report."
        )
    with metrics_path.open(encoding="utf-8") as file:
        results = json.load(file)
    eda = pd.read_csv(eda_path)
    _, test_labels, test_classes = discover_images(ROOT / "test")
    if test_classes != results["class_names"]:
        raise ValueError("Test dataset classes do not match the recorded evaluation.")
    test_counts = pd.Series(test_labels).value_counts().to_dict()
    grouped = eda.groupby("class").agg(
        image_count=("path", "count"),
        median_width=("width", "median"),
        median_height=("height", "median"),
        median_brightness=("mean_brightness", "median"),
        median_dark_fraction=("dark_pixel_fraction", "median"),
        median_bright_fraction=("bright_pixel_fraction", "median"),
    )

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(40, 48, 54)

    title = doc.add_heading("Avocado Image Classification", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(
        "Computer Vision Assignment — Dataset, Methodology, Results and Discussion"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Dataset scope: this repository contains blackspot, canker, fresh, and "
        "grenning image folders. These are the labels actually classified; they "
        "are not the underripe/breaking/ripe/overripe labels in the original brief."
    )

    doc.add_heading("1. Dataset understanding and EDA", level=1)
    doc.add_paragraph(
        f"The supplied training set contains {int(grouped.image_count.sum())} images. "
        f"The separate held-out test set contains {results['test_image_count']} images. "
        "Training data were split once using a stratified 85/15 train/validation "
        "split (842 fit images and 149 validation images); the test folder was "
        "not used for fitting or validation."
    )
    table = doc.add_table(rows=1, cols=7)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [
        "Class",
        "Train",
        "Test",
        "Median W×H",
        "Median brightness",
        "Dark fraction",
        "Bright fraction",
    ]
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    for class_name, row in grouped.iterrows():
        values = [
            class_name,
            str(int(row.image_count)),
            str(test_counts.get(class_name, 0)),
            f"{int(row.median_width)}×{int(row.median_height)}",
            f"{row.median_brightness:.1f}/255",
            f"{row.median_dark_fraction:.2f}",
            f"{row.median_bright_fraction:.2f}",
        ]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    doc.add_paragraph(
        "Lighting and resolution deserve careful bias review. All blackspot, "
        "canker, and grenning training images are 800×800, while fresh images "
        "vary widely in size. Median brightness also differs strongly by class "
        "(approximately 49 for grenning, 95 for canker, 103 for blackspot, "
        "and 155 for fresh on an 8-bit scale). A classifier may exploit capture "
        "or brightness differences rather than only the target visual condition. "
        "The EDA indicators flag distribution shifts for inspection; they do "
        "not establish their cause. The training audit found zero unreadable images."
    )
    add_centered_picture(doc, RESULTS / "eda" / "dataset_overview.png")
    doc.add_paragraph(
        "Per-image resolution and brightness indicators: "
        "report/results/eda/image_analysis.csv"
    )

    doc.add_heading("2. Methodology and model selection", level=1)
    doc.add_heading("Approach 1 — classical CV and SVM", level=2)
    doc.add_paragraph(
        "The classical model combines normalized HSV histograms, CIELAB channel "
        "histograms, and uniform Local Binary Pattern texture histograms. A "
        "standardized feature vector is classified with a class-balanced RBF "
        "Support Vector Machine. This compact baseline exposes color/texture "
        "features directly and provides a useful comparison to a neural model."
    )
    doc.add_heading("Approach 2 — MobileNetV2 transfer learning", level=2)
    doc.add_paragraph(
        "MobileNetV2 starts from ImageNet-pretrained weights, freezes its "
        "convolutional feature backbone, and trains a small four-class head. "
        "Training uses augmentation, class-weighted cross entropy, Adam, and "
        "early stopping against validation loss. The lightweight architecture "
        "is a practical transfer-learning comparison, but production latency "
        "must be measured on the intended device."
    )

    doc.add_heading("3. Held-out model evaluation", level=1)
    doc.add_paragraph(
        f"Both models were scored on the same {results['test_image_count']}-image "
        "test set. Metrics are macro-averaged for precision, recall, and F1 to "
        "give each class equal weight."
    )
    metric_table = doc.add_table(rows=1, cols=5)
    metric_table.style = "Light Shading Accent 1"
    for cell, text in zip(
        metric_table.rows[0].cells,
        ["Model", "Accuracy", "Macro precision", "Macro recall", "Macro F1"],
    ):
        cell.text = text
    for key, label in (
        ("classical_svm", "HSV/CIELAB/LBP + SVM"),
        ("mobilenet_v2", "MobileNetV2"),
    ):
        scores = results["metrics"][key]
        row = metric_table.add_row()
        for cell, value in zip(
            row.cells,
            [
                label,
                f"{scores['accuracy']:.2%}",
                f"{scores['macro_precision']:.2%}",
                f"{scores['macro_recall']:.2%}",
                f"{scores['macro_f1']:.2%}",
            ],
        ):
            cell.text = value

    for key, title_text in (
        ("classical_svm", "Classical SVM confusion matrix"),
        ("mobilenet_v2", "MobileNetV2 confusion matrix"),
    ):
        doc.add_heading(title_text, level=2)
        add_centered_picture(
            doc, RESULTS / f"{key}_confusion_matrix.png", width=5.6
        )
    doc.add_paragraph(
        "The SVM scored 96.97% accuracy and MobileNetV2 95.96% on this test "
        "split; the classical baseline leads by 1.01 percentage points. Both "
        "models classified all fresh and grenning test examples correctly. "
        "SVM canker recall was 86.36% and blackspot precision 88.00%; MobileNetV2 "
        "blackspot recall was 86.36% and canker precision 87.50%. The test set "
        "contains only 22–33 examples per class, so these estimates have "
        "substantial sampling uncertainty and are not production guarantees."
    )

    doc.add_heading("4. Assignment questions", level=1)
    doc.add_heading("4.1 In production, how does your model perform?", level=2)
    doc.add_paragraph(
        "On the supplied test set, the best observed accuracy was 96.97% for "
        "the SVM; MobileNetV2 scored 95.96%. Inference latency was not benchmarked. "
        "These measurements do not establish production performance. Before "
        "deployment, test on independent images from the target camera, lighting, "
        "and capture setup and measure latency on the actual device. This dataset "
        "predicts its four folder labels, not avocado ripeness stage."
    )
    doc.add_heading("4.2 What are the drawbacks?", level=2)
    for text in (
        "The dataset labels do not represent the ripeness stages requested in the original assignment.",
        "Image brightness and resolution vary with class, creating shortcut-learning and dataset-bias risk.",
        "A small test set and possible label ambiguity limit confidence in measured generalization.",
        "A single RGB image captures only visible surface appearance and cannot establish internal firmness or unseen-side condition.",
        "Model confidence scores are not calibrated probabilities.",
    ):
        add_bullet(doc, text)
    doc.add_heading("4.3 What would you improve with more time?", level=2)
    for text in (
        "Acquire the intended ripeness-stage labels or clarify the dataset target, then collect a larger, balanced, representative dataset.",
        "Standardize capture conditions or collect metadata and validate across lighting, camera, background, and source-domain shifts.",
        "Review difficult examples, label disagreements, per-class confusion, and possible train/test near-duplicates.",
        "Evaluate multi-view imaging and non-destructive firmness or spectral sensors if internal ripeness is the actual target.",
        "Benchmark inference and, if appropriate, optimize for the actual deployment hardware while rechecking accuracy.",
    ):
        add_bullet(doc, text)
    doc.add_heading("4.4 Where was the most time spent?", level=2)
    doc.add_paragraph(
        "The project work log did not record the student's personal time allocation, "
        "so percentages cannot be stated truthfully from this run. Complete this "
        "answer with your own approximate hours or percentages for dataset review, "
        "feature/model work, training/debugging, and evaluation/report writing."
    )

    doc.add_heading("5. Reproducibility and submission items", level=1)
    doc.add_paragraph(
        "The notebooks in notebooks/ cover EDA, the classical baseline, "
        "MobileNetV2 training, held-out evaluation, and inference demonstration. "
        "Reproduction commands and dependencies are in README.md. The measured "
        "training ran on CPU for 10 epochs; validation accuracy at the best "
        "validation-loss checkpoint was 100% (validation loss 0.0369)."
    )
    for text in (
        "Public GitHub repository: add the URL after publishing this repository.",
        "Google Colab: upload/open the notebooks and project files, then enable view access for anyone with the link if permitted.",
        "Video walkthrough: record the 30-minute demonstration using the agenda in report/Submission_Guide.md.",
        "Personal time allocation: complete section 4.4 before submission.",
    ):
        add_bullet(doc, text)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Saved report: {OUTPUT}")


if __name__ == "__main__":
    main()
